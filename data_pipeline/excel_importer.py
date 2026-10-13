"""
data_pipeline/excel_importer.py

Reads a teacher's Excel (.xlsx) grade sheet and turns it into validated
StudentRecords. Pure parsing + validation: it never touches the database
(SRP) - persistence is done by GroupImportService.

Expected sheet (first row = headers, one student per row):

    matricula | nombre | parcial1 | parcial2 | parcial3 | parcial4 | asistencia_pct | participacion

Rules (Fail-Fast, validation matrix section 8.0 - "Data Access / ETL" layer):
    - File level, rejected immediately: not a valid .xlsx, empty sheet,
      missing required column, more than MAX_ROWS rows.
    - Row level, the row is rejected (and reported) but the rest continue:
      missing matricula/nombre, non-numeric value, out-of-range value
      (0-100), no grades, gaps between partials, duplicate matricula.
    - At least `parcial1` is required; later partials may be blank while the
      semester is still running.
"""

import io
import re
import unicodedata
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core.exception_handler import ValidationError
from data_access.student_repository import StudentRecord
from data_pipeline.fail_fast_validator import FailFastValidator

MAX_ROWS = 1000
MAX_NAME_LENGTH = 120
MAX_ID_LENGTH = 32

PARTIAL_COLUMNS = ["parcial1", "parcial2", "parcial3", "parcial4"]
TEMPLATE_COLUMNS = ["matricula", "nombre", *PARTIAL_COLUMNS, "asistencia_pct", "participacion"]
REQUIRED_COLUMNS = ["matricula", "nombre", "parcial1", "asistencia_pct", "participacion"]

# Accepted header spellings (after normalisation) -> canonical column name.
_ALIASES = {
    "matricula": "matricula", "id": "matricula", "student_id": "matricula",
    "clave": "matricula", "no_control": "matricula", "numero_de_control": "matricula",
    "nombre": "nombre", "name": "nombre", "alumno": "nombre", "nombre_completo": "nombre",
    "parcial1": "parcial1", "parcial_1": "parcial1", "p1": "parcial1", "calificacion1": "parcial1",
    "parcial2": "parcial2", "parcial_2": "parcial2", "p2": "parcial2", "calificacion2": "parcial2",
    "parcial3": "parcial3", "parcial_3": "parcial3", "p3": "parcial3", "calificacion3": "parcial3",
    "parcial4": "parcial4", "parcial_4": "parcial4", "p4": "parcial4", "calificacion4": "parcial4",
    "asistencia_pct": "asistencia_pct", "asistencia": "asistencia_pct",
    "asistencia_%": "asistencia_pct", "attendance_pct": "asistencia_pct",
    "participacion": "participacion", "participacion_pct": "participacion",
    "participation_score": "participacion", "participation": "participacion",
}


@dataclass
class RejectedRow:
    row: int            # Excel row number (header is row 1)
    student_id: str
    reason: str


@dataclass
class ParsedSheet:
    records: list[StudentRecord] = field(default_factory=list)
    rejected: list[RejectedRow] = field(default_factory=list)
    total_rows: int = 0


def _normalize_header(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    text = re.sub(r"[\s\-\.]+", "_", text.strip().lower())
    return text.strip("_")


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _to_number(value, label: str) -> float:
    """Accepts 85, 85.5, '85', '85,5', '90%'. Rejects anything else."""
    if isinstance(value, bool):
        raise ValueError(f"{label}: valor no numerico")
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("%", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        raise ValueError(f"{label}: '{value}' no es un numero") from None


def build_template() -> bytes:
    """Excel template the teacher can download, fill in and upload back."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Calificaciones"
    ws.append(TEMPLATE_COLUMNS)
    ws.append(["A0001", "Ana Garcia Lopez", 85, 90, 78, 92, 95, 88])
    ws.append(["A0002", "Luis Hernandez Perez", 55, 48, 60, None, 70, 40])

    header_fill = PatternFill("solid", fgColor="1F3A5F")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for column, width in zip("ABCDEFGH", (14, 28, 11, 11, 11, 11, 16, 15)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Instrucciones")
    for line in (
        "Plantilla de importacion - Student Performance Predictor",
        "",
        "1. Una fila por alumno en la hoja 'Calificaciones'. No cambies los encabezados.",
        "2. matricula: identificador unico del alumno (maximo 32 caracteres).",
        "3. nombre: nombre completo del alumno.",
        "4. parcial1..parcial4: calificaciones de 0 a 100. parcial1 es obligatorio;",
        "   los demas pueden quedar vacios si el parcial aun no se evalua (sin huecos).",
        "5. asistencia_pct: porcentaje de asistencia de 0 a 100.",
        "6. participacion: puntaje de participacion de 0 a 100.",
        "7. Las filas de ejemplo (A0001, A0002) pueden borrarse.",
        f"8. Maximo {MAX_ROWS} alumnos por archivo.",
    ):
        info.append([line])
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 90

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def parse_workbook(file_bytes: bytes) -> ParsedSheet:
    """Parse and validate an .xlsx file. Raises ValidationError on file-level problems."""
    if not file_bytes:
        raise ValidationError("empty upload", user_message="El archivo esta vacio.")

    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many different errors
        raise ValidationError(
            f"unreadable workbook: {type(exc).__name__}",
            user_message="No se pudo leer el archivo. Sube un Excel valido (.xlsx).",
        ) from exc

    try:
        ws = wb["Calificaciones"] if "Calificaciones" in wb.sheetnames else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)

        header = next(rows, None)
        if header is None:
            raise ValidationError("no header row", user_message="La hoja no tiene datos.")

        column_index: dict[str, int] = {}
        for position, raw in enumerate(header):
            canonical = _ALIASES.get(_normalize_header(raw))
            if canonical and canonical not in column_index:
                column_index[canonical] = position

        missing = [c for c in REQUIRED_COLUMNS if c not in column_index]
        if missing:
            raise ValidationError(
                f"missing columns: {missing}",
                user_message=(
                    "Faltan columnas obligatorias: " + ", ".join(missing)
                    + ". Descarga la plantilla para ver el formato correcto."
                ),
            )

        result = ParsedSheet()
        seen_ids: set[str] = set()

        for excel_row, row in enumerate(rows, start=2):
            if row is None or all(_is_blank(v) for v in row):
                continue  # silently skip fully empty rows
            result.total_rows += 1
            if result.total_rows > MAX_ROWS:
                raise ValidationError(
                    "too many rows",
                    user_message=f"El archivo supera el maximo de {MAX_ROWS} alumnos.",
                )

            def cell(name: str):
                index = column_index.get(name)
                return row[index] if index is not None and index < len(row) else None

            raw_id = "" if _is_blank(cell("matricula")) else str(cell("matricula")).strip()
            if raw_id.endswith(".0") and raw_id[:-2].isdigit():  # Excel numeric ids -> "123"
                raw_id = raw_id[:-2]

            try:
                record = _build_record(raw_id, cell)
                if record.student_id in seen_ids:
                    raise ValueError(f"matricula '{record.student_id}' repetida en el archivo")
                FailFastValidator.validate(record)
            except ValidationError as exc:
                result.rejected.append(RejectedRow(excel_row, raw_id, str(exc)))
                continue
            except ValueError as exc:
                result.rejected.append(RejectedRow(excel_row, raw_id, str(exc)))
                continue

            seen_ids.add(record.student_id)
            result.records.append(record)

        if result.total_rows == 0:
            raise ValidationError("no data rows", user_message="El archivo no contiene alumnos.")
        return result
    finally:
        wb.close()


def _build_record(raw_id: str, cell) -> StudentRecord:
    if not raw_id:
        raise ValueError("matricula vacia")
    if len(raw_id) > MAX_ID_LENGTH:
        raise ValueError(f"matricula demasiado larga (maximo {MAX_ID_LENGTH} caracteres)")

    name = "" if _is_blank(cell("nombre")) else " ".join(str(cell("nombre")).split())
    if not name:
        raise ValueError("nombre vacio")
    if len(name) > MAX_NAME_LENGTH:
        raise ValueError(f"nombre demasiado largo (maximo {MAX_NAME_LENGTH} caracteres)")

    grades: list[float] = []
    gap_at: str | None = None
    for column in PARTIAL_COLUMNS:
        value = cell(column)
        if _is_blank(value):
            gap_at = gap_at or column
            continue
        if gap_at:
            raise ValueError(f"{gap_at} esta vacio pero {column} tiene calificacion (sin huecos)")
        grades.append(_to_number(value, column))
    if not grades:
        raise ValueError("sin calificaciones: parcial1 es obligatorio")

    for required in ("asistencia_pct", "participacion"):
        if _is_blank(cell(required)):
            raise ValueError(f"{required} vacio")

    return StudentRecord(
        student_id=raw_id,
        name=name,
        grades=grades,
        attendance_pct=_to_number(cell("asistencia_pct"), "asistencia_pct"),
        participation_score=_to_number(cell("participacion"), "participacion"),
    )
