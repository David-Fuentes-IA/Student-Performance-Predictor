"""The fixed evaluation workbook (data/grupo_evaluacion.xlsx): structure, reproducibility,
independence from the training data, and the property the whole of Phase 5 relies on:
a model does NOT score a perfect 1.0 on it."""
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from sklearn.metrics import accuracy_score, confusion_matrix, recall_score

from core.exception_handler import ValidationError
from data_access.synthetic_students import simulate_students
from data_pipeline.excel_importer import TEMPLATE_COLUMNS, parse_workbook
from learning_module.evaluation_data import (
    DEFAULT_PATH,
    LABEL_SHEET,
    load_evaluation_workbook,
)
from learning_module.model_trainer import ModelTrainer, generate_synthetic_training_set
from learning_module.training_pipeline import DATA_SEED, N_SAMPLES
from scripts.build_evaluation_dataset import EVAL_SEED, EVAL_STUDENTS


class WorkbookStructureTests(unittest.TestCase):
    def test_follows_the_import_template(self):
        wb = load_workbook(DEFAULT_PATH)
        self.assertEqual(wb.sheetnames, ["Calificaciones", "Instrucciones", LABEL_SHEET])
        self.assertEqual([c.value for c in wb["Calificaciones"][1]], TEMPLATE_COLUMNS)
        self.assertEqual(wb["Calificaciones"].freeze_panes, "A2")
        self.assertEqual(wb["Calificaciones"].max_row - 1, EVAL_STUDENTS)

    def test_can_be_imported_like_any_real_upload_with_no_rejected_rows(self):
        parsed = parse_workbook(DEFAULT_PATH.read_bytes())
        self.assertEqual((len(parsed.records), len(parsed.rejected)), (EVAL_STUDENTS, 0))
        self.assertTrue(any(len(r.grades) < 4 for r in parsed.records))  # includes unfinished partials

    def test_matches_what_the_simulator_produces(self):
        loaded = load_evaluation_workbook()
        expected = {s.student_id: (s.grades, s.attendance_pct, s.participation_score, s.at_risk, s.shift)
                    for s in simulate_students(EVAL_STUDENTS, seed=EVAL_SEED, id_prefix="E")}
        actual = {r.student_id: (r.grades, r.attendance_pct, r.participation_score,
                                 loaded.labels[r.student_id], loaded.shifts[r.student_id]) for r in loaded.records}
        self.assertEqual(actual, expected)

    def test_outcome_sheet_covers_every_student_with_valid_values(self):
        loaded = load_evaluation_workbook()
        self.assertEqual(set(loaded.labels), {r.student_id for r in loaded.records})
        self.assertEqual(set(loaded.labels.values()), {0, 1})
        self.assertEqual(set(loaded.shifts.values()), {"matutino", "vespertino"})
        share = sum(loaded.labels.values()) / len(loaded.labels)
        self.assertTrue(0.20 <= share <= 0.35, share)


class LoaderRejectionTests(unittest.TestCase):
    def modified_copy(self, modify):
        wb = load_workbook(DEFAULT_PATH)
        modify(wb)
        buffer = io.BytesIO()
        wb.save(buffer)
        tmp = Path(tempfile.mkdtemp()) / "copy.xlsx"
        tmp.write_bytes(buffer.getvalue())
        return tmp

    def test_missing_outcome_sheet(self):
        path = self.modified_copy(lambda wb: wb.remove(wb[LABEL_SHEET]))
        with self.assertRaises(ValidationError):
            load_evaluation_workbook(path)

    def test_outcomes_that_do_not_match_the_students(self):
        path = self.modified_copy(lambda wb: wb[LABEL_SHEET].delete_rows(2))
        with self.assertRaises(ValidationError):
            load_evaluation_workbook(path)


class UnseenDataAndImperfectScoreTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.evaluation = load_evaluation_workbook()
        cls.X, cls.y, cls.shift = cls.evaluation.dataset()
        cls.X_train, cls.y_train = generate_synthetic_training_set(N_SAMPLES, DATA_SEED)

    def test_evaluation_students_are_not_in_the_training_data(self):
        training = {(round(row[0], 2), round(row[2], 1), round(row[3], 1)) for row in self.X_train}
        evaluation = {(round(row[0], 2), round(row[2], 1), round(row[3], 1)) for row in self.X}
        self.assertEqual(training & evaluation, set())

    def test_models_learn_something_but_are_not_perfect(self):
        for name, model in (("random forest", ModelTrainer.train(self.X_train, self.y_train)),
                            ("logistic regression", ModelTrainer.train_logistic_regression(self.X_train, self.y_train))):
            predictions = model.predict(self.X)
            tn, fp, fn, tp = confusion_matrix(self.y, predictions).ravel()
            with self.subTest(name):
                self.assertGreater(accuracy_score(self.y, predictions), 0.70)   # it learned the pattern
                self.assertLess(accuracy_score(self.y, predictions), 0.97)      # but is not a perfect separator
                self.assertLess(recall_score(self.y, predictions), 0.95)
                self.assertGreaterEqual(fn, 5)  # enough false negatives to analyse
                self.assertGreaterEqual(fp, 5)  # enough false positives to analyse


if __name__ == "__main__":
    unittest.main()
