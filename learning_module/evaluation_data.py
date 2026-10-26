"""
learning_module/evaluation_data.py

Reads the fixed evaluation workbook (data/grupo_evaluacion.xlsx).

The workbook follows the import template: sheet `Calificaciones` is parsed exactly like a real
upload (same validation), so the very same file can be imported into the application. A second
sheet, `Resultado_real`, holds what actually happened to each student (the label a model must
predict) and the synthetic subgroup used for the fairness audit. The importer ignores it.

Because these students were simulated with a different seed than the training set, they are
data the model has never seen, which is what validation requires.
"""

import io
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from core.exception_handler import ValidationError
from data_access.student_repository import StudentRecord
from data_pipeline.excel_importer import parse_workbook
from feature_store.feature_manager import FeatureManager

LABEL_SHEET = "Resultado_real"
DEFAULT_PATH = Path(__file__).resolve().parent.parent / "data" / "grupo_evaluacion.xlsx"


@dataclass(frozen=True)
class EvaluationSet:
    records: list[StudentRecord]
    labels: dict[str, int]      # student_id -> 1 if the student ended the term at risk
    shifts: dict[str, str]      # student_id -> synthetic subgroup ("matutino" / "vespertino")

    def dataset(self) -> tuple[list[list[float]], list[int], list[str]]:
        """(X, y, subgroup per row), with features built by the production FeatureManager."""
        vectors = FeatureManager.extract_many(self.records)
        ids = [r.student_id for r in self.records]
        return ([v.as_model_input() for v in vectors], [self.labels[i] for i in ids],
                [self.shifts[i] for i in ids])


def load_evaluation_workbook(path: Path | str = DEFAULT_PATH) -> EvaluationSet:
    content = Path(path).read_bytes()
    parsed = parse_workbook(content)
    if parsed.rejected:
        raise ValidationError(
            f"{len(parsed.rejected)} rows of the evaluation workbook are invalid",
            user_message="The evaluation workbook contains invalid rows.",
        )

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if LABEL_SHEET not in workbook.sheetnames:
            raise ValidationError(f"sheet '{LABEL_SHEET}' missing",
                                  user_message=f"The evaluation workbook has no '{LABEL_SHEET}' sheet.")
        rows = list(workbook[LABEL_SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()

    labels, shifts = {}, {}
    for student_id, at_risk, shift in rows[1:]:
        if student_id is None:
            continue
        labels[str(student_id)] = int(at_risk)
        shifts[str(student_id)] = str(shift)

    ids = {r.student_id for r in parsed.records}
    if ids != set(labels):
        raise ValidationError("labels do not match the students",
                              user_message="'Resultado_real' must list exactly the students of 'Calificaciones'.")
    return EvaluationSet(parsed.records, labels, shifts)
