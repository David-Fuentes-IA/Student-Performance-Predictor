"""
scripts/build_evaluation_dataset.py

Builds data/grupo_evaluacion.xlsx: the fixed, reproducible dataset used to validate the model
(Phase 5) and to demonstrate the application with a realistic class.

    python scripts/build_evaluation_dataset.py              # 240 students, seed 2026
    python scripts/build_evaluation_dataset.py --check      # fail if the committed file differs

It starts from the official import template (sheet `Calificaciones`, same header style, same
`Instrucciones`) and adds a third sheet, `Resultado_real`, with each student's real outcome and
synthetic subgroup. The students are NOT the ones the model is trained on (different seed).
"""

import argparse
import io
import sys
from copy import copy
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from data_access.synthetic_students import SimulatedStudent, simulate_students  # noqa: E402
from data_pipeline.excel_importer import PARTIAL_COLUMNS, build_template  # noqa: E402
from learning_module.evaluation_data import (  # noqa: E402
    DEFAULT_PATH,
    LABEL_SHEET,
    load_evaluation_workbook,
)

EVAL_SEED = 2026
EVAL_STUDENTS = 240
EXTRA_INSTRUCTIONS = (
    "9. ARCHIVO SINTETICO DE EVALUACION: alumnos ficticios (240, semilla 2026), distintos de los de entrenamiento.",
    "10. La hoja 'Resultado_real' guarda el resultado final de cada alumno (en_riesgo: 1 = termino con bajo desempeno)",
    "    y su turno (subgrupo sintetico para la auditoria de fairness). La importacion NO lee esa hoja.",
)


def build_workbook(students: list[SimulatedStudent]) -> bytes:
    wb = load_workbook(io.BytesIO(build_template()))
    grades = wb["Calificaciones"]
    grades.delete_rows(2, grades.max_row)          # drop the two sample rows of the template
    for s in students:
        padded = list(s.grades) + [None] * (len(PARTIAL_COLUMNS) - len(s.grades))
        grades.append([s.student_id, s.name, *padded, s.attendance_pct, s.participation_score])

    instructions = wb["Instrucciones"]
    for line in EXTRA_INSTRUCTIONS:
        instructions.append([line])

    outcome = wb.create_sheet(LABEL_SHEET)
    outcome.append(["matricula", "en_riesgo", "turno"])
    for cell in outcome[1]:
        source = grades["A1"]
        cell.font, cell.fill, cell.alignment = copy(source.font), copy(source.fill), copy(source.alignment)
    for s in students:
        outcome.append([s.student_id, s.at_risk, s.shift])
    for column, width in zip("ABC", (14, 12, 14), strict=True):
        outcome.column_dimensions[column].width = width
    outcome.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the fixed evaluation workbook.")
    parser.add_argument("--out", type=Path, default=DEFAULT_PATH)
    parser.add_argument("--check", action="store_true", help="verify the existing file instead of writing")
    args = parser.parse_args()

    students = simulate_students(EVAL_STUDENTS, seed=EVAL_SEED, id_prefix="E")
    if args.check:
        loaded = load_evaluation_workbook(args.out)
        expected = {s.student_id: (s.grades, s.attendance_pct, s.participation_score, s.at_risk, s.shift)
                    for s in students}
        actual = {r.student_id: (r.grades, r.attendance_pct, r.participation_score,
                                 loaded.labels[r.student_id], loaded.shifts[r.student_id]) for r in loaded.records}
        if expected != actual:
            raise SystemExit(f"{args.out} does not match the simulator output: rebuild it.")
        print(f"{args.out} matches the simulator ({len(students)} students).")
        return

    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_bytes(build_workbook(students))
    at_risk = sum(s.at_risk for s in students)
    print(f"Wrote {args.out}: {len(students)} students, {at_risk} at risk ({100 * at_risk / len(students):.1f} %).")


if __name__ == "__main__":
    main()
